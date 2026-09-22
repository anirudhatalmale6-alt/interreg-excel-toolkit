#!/usr/bin/env python3
"""Compare successive MIS technical-sheet exports for EPIRUSMEDEYE.

WHY NOT THE GEEKSFORGEEKS APPROACH

That article loads two sheets into pandas and compares them cell by cell.  On
these files that would be worse than useless, because the three exports have
127, 120 and 120 rows.  Every row after the first difference is shifted, so a
positional comparison reports hundreds of false differences and buries the two
or three real ones.

So this compares on MEANING, not position: each amount is keyed by
(beneficiary code, deliverable, cost category).  Rows can move, blocks can be
reordered, a deliverable can be inserted or deleted, and the diff still says
exactly what changed.

TWO THINGS THE COLUMN LAYOUT TRIES TO TRICK YOU WITH

1. Row 8 of each partner block carries MIS budget-line codes ("1 | B.24",
   "7 | B.30").  They look like a column map.  They are not - they appear in
   DIFFERENT columns for different partners, because the export only writes
   codes for the lines that partner actually uses, and misaligns them.  The
   reliable map is the row-6 TITLE row, which is identical in every block.

2. So the column mapping here is asserted rather than assumed: for every
   deliverable row in every file, the six category columns must add up to the
   row total in column 10.  If the mapping were wrong that check fails
   immediately instead of producing a confident wrong answer.

The closure check at the end matters just as much: the sum of every individual
change must equal the change in the grand total.  If it does not, the diff has
missed something, and a diff that quietly misses a change is the only really
dangerous outcome here.
"""
import sys
from pathlib import Path

import openpyxl

# WHERE THE FILES ARE, AND WHY IT IS WRITTEN THIS WAY
#
# This used to say BASE = "/var/lib/freelancer/projects/40536008" - a path on
# MY machine.  On yours that folder does not exist, so the script raised
# FileNotFoundError no matter how the run configuration was set.  That was my
# mistake, not a setup problem.
#
# Path(__file__) is the script's own location, and .parent is the folder it
# lives in.  So the rule is simply: keep the three .xlsx files in the same
# folder as this .py file and it works on Windows, Linux or anywhere else with
# nothing to edit.
#
# It also makes PyCharm's "Working directory" setting irrelevant here.  A
# relative filename like "1.0-interreg_report26375.xlsx" is resolved against
# the working directory, which is why that field can silently break a script.
# An absolute path built from __file__ cannot.
BASE = Path(__file__).resolve().parent

FILES = [
    ("v1.0", "1.0-interreg_report26375.xlsx"),
    ("v1.1", "1.1-interreg_report96606.xlsx"),
    ("v2",   "2-interreg_report129557.xlsx"),
]

# column -> cost category, taken from the row-6 title row (stable across blocks)
CATS = {
    3: "Staff costs",
    4: "Overheads",
    5: "Travel & accommodation",
    7: "External expertise & services",
    8: "Equipment",
    9: "Investments / infrastructure",
}
TOTAL_COL = 10


def parse(path):
    """-> (cells, partners, grand)

    cells    {(ben_code, deliverable, category): amount}
    partners {ben_code: {"name":..., "total": declared ΓΕΝΙΚΟ ΣΥΝΟΛΟ}}
    grand    sum of the declared partner totals
    """
    ws = openpyxl.load_workbook(path, data_only=True).active
    cells, partners = {}, {}
    ben = None

    for r in range(1, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value
        a = str(a).strip() if a is not None else ""

        # a partner block header, e.g. "Κύριος Δικαιούχος 1090210 (...)" or "P2 1083528 (...)"
        digits = [t for t in a.replace("(", " ").split() if t.isdigit() and len(t) >= 6]
        if digits and "(" in a:
            ben = digits[0]
            partners.setdefault(ben, {"name": a, "total": None})
            continue

        if ben is None:
            continue

        # the declared project total for this partner
        if a.startswith("ΓΕΝΙΚΟ ΣΥΝΟΛΟ"):
            v = ws.cell(row=r, column=TOTAL_COL).value
            if isinstance(v, (int, float)):
                partners[ben]["total"] = float(v)
            continue

        # a deliverable row
        if not a.startswith("Παραδοτέο"):
            continue
        deliv = " ".join(a.split())                      # "Παραδοτέο 1. 2" -> normalised
        deliv = deliv.replace("Παραδοτέο ", "D.").replace(". ", ".")

        rowsum = 0.0
        for col, cat in CATS.items():
            v = ws.cell(row=r, column=col).value
            amt = float(v) if isinstance(v, (int, float)) else 0.0
            rowsum += amt
            if amt:
                cells[(ben, deliv, cat)] = amt

        # PROVE the column mapping instead of trusting it
        declared = ws.cell(row=r, column=TOTAL_COL).value
        if isinstance(declared, (int, float)):
            assert abs(rowsum - float(declared)) < 0.02, (
                f"{path} row {r} {deliv}: categories sum to {rowsum:,.2f} but the row "
                f"total says {float(declared):,.2f} — the column mapping is wrong"
            )

    grand = sum(p["total"] for p in partners.values() if p["total"] is not None)
    return cells, partners, grand


def diff(a, b):
    """Changes from a to b, keyed on meaning so row moves are invisible."""
    out = []
    for k in sorted(set(a) | set(b), key=lambda t: (t[0], t[1], t[2])):
        x, y = a.get(k, 0.0), b.get(k, 0.0)
        if abs(x - y) > 0.005:
            out.append((k, x, y, y - x))
    return out


def main():
    loaded = []
    print("PARSED")
    for label, name in FILES:
        path = BASE / name          # Path joins with / on every platform
        if not path.exists():
            # say WHICH file and WHERE it looked - "FileNotFoundError" alone
            # sends you hunting through the run configuration for an hour
            print(f"Cannot find {name}\n  looked in: {BASE}\n"
                  f"  put the three .xlsx files in that folder, "
                  f"beside this script.")
            return 1
        cells, partners, grand = parse(path)
        loaded.append((label, name, cells, partners, grand))
        print(f"  {label:5s} {len(partners)} partners, {len(cells):4d} non-zero amounts, "
              f"grand total {grand:>14,.2f}   [{name}]")

    # ---- partner totals side by side
    print("\nPARTNER TOTALS BY VERSION")
    codes = []
    for _, _, _, partners, _ in loaded:
        for c in partners:
            if c not in codes:
                codes.append(c)
    hdr = "  {:<10}".format("code") + "".join(f"{l:>16}" for l, *_ in loaded) + "   name"
    print(hdr)
    for c in codes:
        line = f"  {c:<10}"
        for _, _, _, partners, _ in loaded:
            t = partners.get(c, {}).get("total")
            line += f"{t:>16,.2f}" if t is not None else f"{'—':>16}"
        nm = next((p[c]["name"] for _, _, _, p, _ in loaded if c in p), "")
        print(line + "   " + nm[:56])
    line = "  {:<10}".format("TOTAL")
    for _, _, _, _, grand in loaded:
        line += f"{grand:>16,.2f}"
    print(line)

    # ---- consecutive diffs
    for i in range(len(loaded) - 1):
        la, _, ca, pa, ga = loaded[i]
        lb, _, cb, pb, gb = loaded[i + 1]
        ch = diff(ca, cb)
        print(f"\n{'='*104}\n{la} -> {lb}     grand total "
              f"{ga:,.2f} -> {gb:,.2f}   ({gb - ga:+,.2f})\n{'='*104}")
        if not ch:
            print("  no change to any amount")
        for (ben, deliv, cat), x, y, d in ch:
            nm = pb.get(ben, pa.get(ben, {})).get("name", ben)
            short = nm.split("(")[-1].rstrip(")")[:30] if "(" in nm else nm[:30]
            kind = "ADDED  " if x == 0 else ("REMOVED" if y == 0 else "CHANGED")
            print(f"  {kind} {ben} {short:<32} {deliv:<9} {cat:<30} "
                  f"{x:>12,.2f} -> {y:>12,.2f}  {d:+12,.2f}")

        # CLOSURE: the changes must account for the whole movement in the total
        moved = sum(d for *_, d in ch)
        print(f"\n  sum of the {len(ch)} change(s): {moved:+,.2f}   "
              f"grand total moved: {gb - ga:+,.2f}")
        if abs(moved - (gb - ga)) > 0.02:
            print(f"  ** MISMATCH of {moved - (gb - ga):+,.2f} — the diff is incomplete, "
                  f"or a partner total was restated without its deliverables changing")
        else:
            print("  reconciled — every cent of the movement is explained above")

    # ---- first vs last
    la, _, ca, pa, ga = loaded[0]
    lb, _, cb, pb, gb = loaded[-1]
    ch = diff(ca, cb)
    print(f"\n{'='*104}\nNET {la} -> {lb}   {ga:,.2f} -> {gb:,.2f}   ({gb - ga:+,.2f}), "
          f"{len(ch)} amount(s) differ\n{'='*104}")
    for (ben, deliv, cat), x, y, d in ch:
        print(f"  {ben} {deliv:<9} {cat:<30} {x:>12,.2f} -> {y:>12,.2f}  {d:+12,.2f}")
    return loaded


if __name__ == "__main__":
    sys.exit(0 if main() else 0)
