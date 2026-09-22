"""
LESSON 2 - your own script, corrected, with the reasons.

This is YOUR OPENPYXL.py from the screenshot. I have not rewritten it into
something fancier, because the point is for you to see what was wrong with the
thing you actually typed. Your version is at the bottom, commented out, so you
can compare line by line.

Five corrections. The last one is the one that can destroy a real file.
"""

import os

from openpyxl import load_workbook

# ---------------------------------------------------------------- CORRECTION 1
# You had:   from openpyxl.workbook import workbook
#
# Two problems. Lowercase 'workbook' is the MODULE; the class is 'Workbook'
# with a capital W. And you do not need it at all - Workbook() is for CREATING
# a new file from nothing. You are OPENING an existing one, which is
# load_workbook. So the whole line goes.
#
# Rule of thumb: import only what you actually use. An unused import that is
# also misspelled gives you a red underline that tells you nothing useful.


# ---------------------------------------------------------------- CORRECTION 2
# "Do I need to put the files in the directory of the script?"
#
# Either put them in the same folder, or give the full path. There is no third
# option - Python does not go looking.
#
# But be careful: 'the same folder as the script' is not always where Python
# looks. Python looks in the CURRENT WORKING DIRECTORY, which is usually the
# folder your editor opened, not the folder the .py file sits in. When they
# differ you get FileNotFoundError on a file you can see with your own eyes.
#
# So print it. Every time a file is 'missing', this one line tells you why:
print("Python is looking in:", os.getcwd())
print("This script lives in:", os.path.dirname(os.path.abspath(__file__)))

# The robust version - build the path relative to the SCRIPT, not to wherever
# the editor happened to start. Then it works no matter how you run it:
HERE = os.path.dirname(os.path.abspath(__file__))
FILE = os.path.join(HERE, "Tutorial.xlsx")

# A full Windows path also works, and note the r before the quote. Without it,
# \t and \n inside a path become a tab and a newline:
#   FILE = r"D:\PROJECTS\PYTHON\EXCEL-OPENPYXL\Tutorial.xlsx"

wb = load_workbook(filename=FILE)
ws = wb.active


# ---------------------------------------------------------------- CORRECTION 3
# You had:
#     col_a   = ws['A']
#     row_two = ws['B']      <-- this is COLUMN B, not row two
#     row_three = ws['A2']   <-- this is ONE CELL, not row three
#
# The three forms are genuinely different, and mixing them up is the single
# most common openpyxl confusion:
#
#     ws['A']    -> a whole COLUMN   (letters are columns)
#     ws[2]      -> a whole ROW      (bare numbers are rows)
#     ws['A2']   -> a single CELL    (letter + number)
#
col_a = ws["A"]          # column A
col_b = ws["B"]          # column B  (what you called row_two)
row_two = ws[2]          # row 2     (what you meant)
cell_a2 = ws["A2"]       # the single cell A2

print("column A has", len(col_a), "cells")
print("row 2 has", len(row_two), "cells")
print("A2 contains:", cell_a2.value)      # .value - the cell is an object

# Note .value above. ws['A2'] gives you a Cell OBJECT, not the contents.
# Printing it directly shows <Cell 'Sheet'.A2>, which confuses everyone once.


# ---------------------------------------------------------------- CORRECTION 4
# You had:   for row in ws.iter_rows(min_row=2, max_row5, min_col=1, ...)
#                                                ^^^^^^^^
# Missing '='. That is a SyntaxError, so the file would not run at all - not
# one line of it. Python checks the whole file before running anything, which
# is why you get nothing rather than partial output.
for row in ws.iter_rows(min_row=2, max_row=5, min_col=1, max_col=4,
                        values_only=True):
    print(row)          # the whole row as a tuple
    # row[0] is just the first column. Printing the whole tuple while you are
    # learning shows you the shape of what you are getting back.

for col in ws.iter_cols(min_row=1, max_row=5, min_col=1, max_col=4,
                        values_only=True):
    print(col)

# values_only=True gives you the plain values. Leave it out and you get Cell
# objects, which is what you want if you need .row, .column or .number_format.


# ---------------------------------------------------------------- CORRECTION 5
# THE DANGEROUS ONE.
#
# You had:   wb.save('Tutorial.xlsx')
#
# Two separate problems.
#
# (a) You changed nothing, so that line only risks the file without gaining
#     anything. Saving is not "closing". You do not need to save to read.
#
# (b) The one that matters. openpyxl does NOT preserve formulas when you load
#     with data_only=True and then save. It replaces every formula with the
#     last value Excel cached. Load an Annex 3 that way, save it, and you have
#     permanently turned a live budget workbook into flat numbers. Nothing
#     warns you. It looks identical until someone edits a cell and discovers
#     nothing recalculates.
#
#     This matters for you specifically, because in lesson 1 I told you to use
#     data_only=True to read the Annex 3. That is correct for READING. Just
#     never save that same workbook object.
#
# Two habits that make this impossible to get wrong:
#     1. Never save over the file you opened. Write to a new name.
#     2. If you loaded with data_only=True, treat the workbook as read-only.
#
# OUT = os.path.join(HERE, "Tutorial_output.xlsx")
# wb.save(OUT)          # a NEW file - the original is untouched

print("\nDone. Nothing was saved, so Tutorial.xlsx is untouched.")


# ---------------------------------------------------------------------------
# YOUR ORIGINAL, for comparison:
#
# from openpyxl.workbook import workbook        # 1. unused + wrong case
# from openpyxl import load_workbook
#
# wb=load_workbook(filename='Tutorial.xlsx')    # 2. bare filename
# ws=wb.active
# col_a=ws['A']
# row_two=ws['B']                               # 3. that is column B
# row_three=ws['A2']                            # 3. that is one cell
#
# for row in ws.iter_rows(min_row=2, max_row5, min_col=1, max_col=4,
#                         values_only=True):    # 4. SyntaxError, missing =
#     print(row[0])
#
# for col in ws.iter_cols(min_row=1, min_col=1, max_col=4, values_only=True):
#         print(col[0])
#
# wb.save('Tutorial.xlsx')                      # 5. saves over the source
